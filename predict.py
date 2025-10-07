import os
import serial
import threading
import sys
from collections import deque, defaultdict
import pandas as pd
import numpy as np
from dash import Dash, dcc, html
from dash.dependencies import Input, Output, State
from datetime import datetime
import csv
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

# ---------------- CONFIG ----------------
SERIAL_PORT = "COM6"
BAUDRATE = 115200
MAX_BUFFER_LINES = 500

serial_lock = threading.Lock()

# Buffers per sensor ID
serial_buffer = defaultdict(lambda: deque(maxlen=MAX_BUFFER_LINES))

start_time = datetime.now()

# ---------------- CSV LOGGING ----------------
timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
csv_file = f"{timestamp_str}.csv"
annotation_file = f"{timestamp_str}.txt"

csv_columns = ["id","index","millis","gas_index","mes_index","temperature","pressure","humidity","gas_resistance","status","date_time"]

# Initialize CSV
with open(csv_file, mode="w", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=csv_columns)
    writer.writeheader()

def try_serial():
    try:
        ser = serial.Serial(SERIAL_PORT, BAUDRATE, timeout=0.5)
        return ser
    except Exception as e:
        print(f"[serial_reader] Error opening {SERIAL_PORT}: {e}")  

def serial_reader():
    ser = try_serial()
    while ser == None:
        ser = try_serial()

    while True:
        try:
            line_bytes = ser.readline()
            if not line_bytes:
                continue
            line = line_bytes.decode("utf-8", errors="ignore").strip()
            if not line:
                continue

            parts = line.split(",")
            if len(parts) < 10:
                continue
            sensor_id = parts[0]
            gas_resistance = float(parts[8])
            temperature = float(parts[5])
            pressure = float(parts[6])
            humidity = float(parts[7])
            status = parts[9]
            gas_index = int(parts[3])
            mes_index = int(parts[4])
            index = int(parts[1])
            millis = int(parts[2])

            # Append to serial buffer for plotting
            with serial_lock:
                # serial_buffer[sensor_id].append({"millis": millis, "gas_resistance": gas_resistance})
                serial_buffer[sensor_id].append({
                    "timestamp": datetime.now(),
                    "gas_resistance": gas_resistance,
                })

            # ---------------- LOG TO CSV ----------------
            row = {
                "id": sensor_id,
                "index": index,
                "millis": millis,
                "gas_index": gas_index,
                "mes_index": mes_index,
                "temperature": temperature,
                "pressure": pressure,
                "humidity": humidity,
                "gas_resistance": gas_resistance,
                "status": status,
                "date_time" : datetime.now().strftime("%d-%m-%Y-%H:%M:%S%f")[:-2]
            }

            csv_columns_full = list(row.keys())
            with open(csv_file, mode="a", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=csv_columns_full)
                writer.writerow(row)
            
            #write_excel_row(row, csv_columns_full, f"{timestamp_str}.xlsx")
            # ---------------- PRINT OUTPUT ----------------
            #print(row)

        except Exception as e:
            print("[serial_reader] Error:", e)
            ser = try_serial()
            while ser == None:
                ser = try_serial()



def write_excel_row(row: dict, csv_columns_full: list, excel_file: str):
    """
    Appends a row (dict) to an Excel file (.xlsx).
    If the file doesn't exist, creates it with the given column headers.
    Keeps decimal numbers as numeric values.
    """

    if not excel_file.lower().endswith(".xlsx"):
        raise ValueError("File must have .xlsx extension")

    try:
        # Create new workbook if it doesn't exist
        if not os.path.exists(excel_file):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Data"

            # Write header row
            ws.append(csv_columns_full)

            wb.save(excel_file)
            wb.close()  # Close immediately

        # Load existing workbook
        wb = openpyxl.load_workbook(excel_file)
        ws = wb.active

        # Check headers match
        existing_headers = [cell.value for cell in ws[1]]
        if existing_headers != csv_columns_full:
            raise ValueError("Column mismatch: Excel headers differ from csv_columns_full")

        # Ensure numeric values are numbers (not strings)
        row_values = []
        for col in csv_columns_full:
            value = row.get(col, None)
            # Convert numeric strings to float if possible
            if isinstance(value, str):
                try:
                    if '.' in value or value.isdigit():
                        value = float(value)
                except:
                    pass
            row_values.append(value)

        # Append the row
        ws.append(row_values)

        # Optional: auto-adjust column widths
        for i, col_name in enumerate(csv_columns_full, start=1):
            max_length = max(len(str(col_name)),
                             *(len(str(cell.value)) for cell in ws[get_column_letter(i)] if cell.value))
            ws.column_dimensions[get_column_letter(i)].width = max_length + 2

        wb.save(excel_file)
        wb.close()

    except Exception as e:
        print("Error writing Excel row:", e)
   



def save_annotations(value):
    """Writes a value to a text file; appends if it exists, creates if not."""
    mode = "a" if os.path.exists(annotation_file) else "w"
    with open(annotation_file, mode, encoding="utf-8") as file:
        file.write(f"{value}\n")

# ---------------- DASH APP ----------------
app = Dash(__name__)

app.layout = html.Div([
    html.H2("BME688 Leonardo's Data", style={"color":"white","textAlign":"center"}),
    # 🔹 Input + Button to store data
    html.Div([
        dcc.Input(
            id="user-input",
            type="text",
            placeholder="Digite a descrição",
            style={"marginRight": "10px", "width": "300px", "height" : "40px"}
        ),
        html.Button(
            "Salvar",
            id="save-btn",
            n_clicks=0,
            disabled=True,
            style={"backgroundColor": "#008CBA", "color": "white", "padding": "8px 16px", "border": "none", "borderRadius": "6px"}
        ),
        html.Div(id="saved-output", style={"color": "white", "marginTop": "10px"}),
        dcc.Store(id="stored-variable")
    ], style={"marginBottom": "30px", "textAlign": "center"}),
    html.Div(id="sensor-predictions", style={"display":"grid","gridTemplateColumns":"repeat(4, 1fr)","gap":"10px"}),
    dcc.Graph(id="live-graph", style={"height":"80vh"}),
    dcc.Interval(id="interval-refresh", interval=1000, n_intervals=0),
], style={"backgroundColor":"#111","padding":"20px"})


# ---------------- USER INPUT HANDLER ----------------
@app.callback(
    Output("save-btn", "disabled"),
    Input("user-input", "value")
)
def toggle_button_state(value):
    """Enable button only when input is not empty."""
    if value and value.strip():
        return False
    return True


@app.callback(
    [Output("stored-variable", "data"),
     Output("saved-output", "children")],
    Input("save-btn", "n_clicks"),
    State("user-input", "value"),
    prevent_initial_call=True
)
def save_user_input(n_clicks, value):
    """Save the input value when button is clicked."""
    if value:
        save_annotations(f"{datetime.now()} - {value}")
        return value, f"✅ Descrição salva: {value}"
    return None, "No data saved"

@app.callback(
    [Output("sensor-predictions","children"),
     Output("live-graph","figure")],
    [Input("interval-refresh","n_intervals")],
)
def update_dashboard(n):
    import plotly.graph_objects as go
    from plotly.subplots import make_subplots

    with serial_lock:
        sensor_ids = list(serial_buffer.keys())
        #To avoid cut sensor_id and get the correct sensor_id list
        for sensor in sensor_ids:
            if(len(sensor) < 9):
                print("------------------------------------------------------------")
                print("WARNING - THERE IS A SENSOR_ID CUT, LETS GET THE KEYS AGAIN")
                print("------------------------------------------------------------")
                sensor_ids = list(serial_buffer.keys())
        dfs = {sid: pd.DataFrame(list(serial_buffer[sid])) for sid in sensor_ids}

    if not sensor_ids:
        return ["No sensors"], go.Figure()

    cards = []


    rows = int(np.ceil(len(sensor_ids)/2))
    cols = 2 if len(sensor_ids) > 1 else 1
    fig = make_subplots(rows=rows, cols=cols, subplot_titles=[f"Sensor {sid}" for sid in sensor_ids])
    for idx, sid in enumerate(sensor_ids):
        df = dfs[sid]
        if df.empty:
            continue

        row = idx // cols + 1
        col = idx % cols + 1

        # inside your callback, for each df:
        df = dfs[sid].copy()
        if df.empty:
            continue

        # make sure timestamp column is datetime
        df['timestamp'] = pd.to_datetime(df['timestamp'])

        # optional: show relative seconds from start_time
        df['elapsed_s'] = (df['timestamp'] - start_time).dt.total_seconds()

        # line=dict(color="#00FFFF") -> TO CHANGE THE LINES COLOR INCLUDING THIS PROPERTY
        fig.add_trace(go.Scatter(x=df['timestamp'], y=df['gas_resistance'], mode='lines', name=f'Sensor ID {sid}'), row=row, col=col)
        

    fig.update_layout(template="plotly_dark", hovermode="x unified", height=1200)
    fig.update_yaxes(title_text="(Ω, log)", type="log")

    return cards, fig

# ---------------- MAIN ----------------
if __name__ == "__main__":
    threading.Thread(target=serial_reader, daemon=True).start()
    app.run(debug=True, use_reloader=False)
